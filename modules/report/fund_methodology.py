"""
Fund methodology export: shows, for one fund on one recalculation date (its inception by
default), how its holdings were derived — each constituent ETF's holdings and active-weight
calculation, the benchmark(s) it was compared against, the best ideas each ETF produced, and
how those were filtered, ranked and weighted into the fund.

Reads only: the per-ETF calculation comes from the fund_analysis snapshot written by
funds_update.activate_fund on each recalculation (live and sim). The fund-level selection is
re-run with model_fund.generate — a pure function — on the same stored best ideas, and
compared with the stored fund_holding rows (matches_stored).
"""
import os
import re
import json
from datetime import date, timedelta
import pandas as pd
from openpyxl.utils import get_column_letter
from modules.const import LARGE_CAP_THRESHOLD
from modules.calc import model_fund
from modules.cron import best_ideas_generator as big
from modules.cron import funds_update
from modules.object import (
    benchmark, best_idea, fund, fund_analysis, fund_holding, fund_holding_change,
    provider, provider_etf, ticker,
)
from modules.object.provider_etf_holding import DUP_PRICE_TOLERANCE, aggregate_holdings, fetch_latest_holdings_for_etf, implied_price

OUTPUT_ROOT = os.path.join('.output', 'methodology')

PCT = '0.00%'
MONEY = '#,##0'
NUM = '#,##0.00'


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', '-', name).strip()


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r'[\[\]:*?/\\]', '-', name)[:31]
    candidate, n = base, 2
    while candidate in used:
        suffix = f" ({n})"
        candidate, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(candidate)
    return candidate


def _write_sheet(
    writer: pd.ExcelWriter,
    sheet: str,
    df: pd.DataFrame,
    formats: dict[str, str] | None = None,
    header: list[tuple[str, object]] | None = None,
) -> None:
    """Writes df (optionally under key/value header lines), freezes its header row, applies
    number formats per column and sizes columns to their content."""
    start = len(header) + 1 if header else 0
    df.to_excel(writer, sheet_name=sheet, index=False, startrow=start)
    ws = writer.sheets[sheet]
    for i, (k, v) in enumerate(header or [], start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    ws.freeze_panes = ws.cell(row=start + 2, column=1)
    for col_idx, col in enumerate(df.columns, start=1):
        fmt = (formats or {}).get(col)
        if fmt:
            for row in range(start + 2, start + 2 + len(df)):
                ws.cell(row=row, column=col_idx).number_format = fmt
        values = [str(col)] + [str(v) for v in df[col].head(200).tolist()]
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(v) for v in values) + 2, 60)


class _Context:
    """Everything loaded once and shared by the individual file writers."""

    def __init__(self, fund_id: int, as_of_date: date | None):
        f = fund.fetch_by_id(fund_id)
        if f is None:
            raise RuntimeError(f"Fund not found: fund_id={fund_id}")
        self.fund = f
        self.strategy = model_fund.getStrategyFromJson(f.strategy)
        self.mode = self.strategy.benchmark

        inception = fund_holding.fetch_first_holding_date(fund_id)
        if inception is None:
            raise RuntimeError(f"Fund {fund_id} has no holdings yet — nothing to explain.")
        self.inception = inception
        self.as_of_date = as_of_date or inception

        self.analysis = fund_analysis.fetch_for_fund_date(fund_id, self.as_of_date)
        if not self.analysis:
            raise RuntimeError(
                f"No fund_analysis snapshot for fund {fund_id} on {self.as_of_date}. It's written on each fund "
                f"recalculation — rerun the sim (scripts/sim_fund.py) for this fund, or wait for the next live recalc."
            )

        stored = fund_holding.fetch_funds_holdings(fund_id, self.as_of_date)
        self.stored_holdings = [h for h in stored if h.holding_date == self.as_of_date]
        self.buys = {c.ticker_id: c for c in fund_holding_change.fetch_for_date(fund_id, self.as_of_date) if c.direction == 'buy'}

        # Re-run the fund's selection exactly as funds_update.activate_fund did.
        all_df, self.mc_map = funds_update.build_shared_context(self.as_of_date)
        self.ideas_df = all_df[all_df['benchmark_mode'] == self.mode]
        self.previous = fund_holding.fetch_funds_holdings(fund_id, self.as_of_date - timedelta(days=1))
        self.held_ids = {p.ticker_id for p in self.previous}
        self.recomputed = model_fund.generate(
            today=self.as_of_date, fund=model_fund.to_fund_protocol(f),
            previous_holdings=[model_fund.FundHolding(**vars(p)) for p in self.previous],
            all_best_ideas_df=self.ideas_df, mc_map=self.mc_map,
        )
        _ideal, self.candidates = model_fund._fetch_and_select_by_region(self.strategy, self.ideas_df, self.held_ids)

        # ETFs, keyed by id, with the holdings date the fund used
        self.etf_dates = {a.provider_etf_id: a.holding_date for a in self.analysis}
        self.etfs = {e: provider_etf.fetch_by_id(e) for e in self.etf_dates}
        self.providers = {p.id: p for p in provider.fetch_by_ids(list({e.provider_id for e in self.etfs.values()}))}
        bm_ids = list({a.benchmark_id for a in self.analysis if a.benchmark_id})
        self.benchmarks = {b.id: b for b in benchmark.fetch_by_ids(bm_ids)} if bm_ids else {}

        # Raw holding lines per ETF (for listing-level detail and quarantined lines)
        self.raw_lines = {
            e: fetch_latest_holdings_for_etf(e, big.LOOK_BACK_WINDOW, up_to_date=d) for e, d in self.etf_dates.items()
        }

        ids = {a.ticker_id for a in self.analysis}
        ids |= {h.ticker_id for lines in self.raw_lines.values() for h in lines if h.ticker_id}
        ids |= {int(t) for t in self.ideas_df['ticker_id']}
        ids |= {h.ticker_id for h in self.stored_holdings} | {h.ticker_id for h in self.recomputed.holdings}
        self.tickers = {t.id: t for t in ticker.fetch_by_ids(list(ids))}

    def sym(self, ticker_id: int | None) -> str:
        t = self.tickers.get(ticker_id) if ticker_id else None
        return t.symbol if t else str(ticker_id or '')

    def tname(self, ticker_id: int | None) -> str:
        t = self.tickers.get(ticker_id) if ticker_id else None
        return (t.name or '') if t else ''

    def exchange(self, ticker_id: int | None) -> str:
        t = self.tickers.get(ticker_id) if ticker_id else None
        return (t.exchange or '') if t else ''

    def provider_name(self, etf_id: int) -> str:
        e = self.etfs.get(etf_id)
        p = self.providers.get(e.provider_id) if e else None
        return p.name if p else ''

    def etf_label(self, etf_id: int | None) -> str:
        e = self.etfs.get(etf_id) if etf_id else None
        if e is None and etf_id:
            e = provider_etf.fetch_by_id(etf_id)
            self.etfs[etf_id] = e
        return f"{e.name} ({etf_id})" if e else str(etf_id or '')

    def company_of(self, ticker_id: int) -> int:
        t = self.tickers.get(ticker_id)
        return (t.master_ticker_id or ticker_id) if t else ticker_id


# ── ETF files ────────────────────────────────────────────────────────────────

def _write_etf_files(ctx: _Context, out_dir: str) -> list[str]:
    etf_dir = os.path.join(out_dir, 'etfs')
    os.makedirs(etf_dir, exist_ok=True)
    written = []
    by_etf: dict[int, list[fund_analysis.FundAnalysis]] = {}
    for a in ctx.analysis:
        by_etf.setdefault(a.provider_etf_id, []).append(a)

    for etf_id, rows in by_etf.items():
        pe = ctx.etfs[etf_id]
        holdings, quarantined = aggregate_holdings(ctx.raw_lines[etf_id])
        line_counts: dict[int, int] = {}
        for h in ctx.raw_lines[etf_id]:
            line_counts[h.ticker_id] = line_counts.get(h.ticker_id, 0) + 1

        # Listing-level detail, grouped under the company the calculation used
        company_ids = {a.ticker_id for a in rows}
        listings = pd.DataFrame([{
            'Company': ctx.sym(ctx.company_of(h.ticker_id) if ctx.company_of(h.ticker_id) in company_ids else h.ticker_id),
            'Listing': ctx.sym(h.ticker_id),
            'Exchange': ctx.tickers[h.ticker_id].exchange if h.ticker_id in ctx.tickers else '',
            'Name': ctx.tname(h.ticker_id),
            'Shares': h.shares,
            'Market value': h.market_value,
            'Provider weight': h.weight,
            'Lines in file': line_counts.get(h.ticker_id, 1),
        } for h in holdings if h.ticker_id])
        per_company = listings.groupby('Company').agg(
            listings=('Listing', lambda s: ', '.join(s)),
            shares=('Shares', 'sum'), market_value=('Market value', 'sum'),
        ) if not listings.empty else pd.DataFrame()

        company = pd.DataFrame([{
            'Rank': a.ranking,
            'Symbol': ctx.sym(a.ticker_id),
            'Name': ctx.tname(a.ticker_id),
            'Country': ctx.tickers[a.ticker_id].country if a.ticker_id in ctx.tickers else '',
            'Region': ctx.tickers[a.ticker_id].region if a.ticker_id in ctx.tickers else '',
            'Listings held': per_company['listings'].get(ctx.sym(a.ticker_id), '') if not per_company.empty else '',
            'Master used': 'Y' if a.master_used else 'N',
            'Market value': per_company['market_value'].get(ctx.sym(a.ticker_id)) if not per_company.empty else None,
            'Market cap used': a.market_cap,
            'ETF weight': a.etf_weight,
            'Benchmark weight': a.benchmark_weight,
            'Delta': a.delta,
            'Note': a.note,
        } for a in rows])
        company = company.sort_values(['Rank', 'Delta'], ascending=[True, False], na_position='last')

        bm = ctx.benchmarks.get(rows[0].benchmark_id) if rows[0].benchmark_id else None
        p = ctx.providers.get(pe.provider_id)
        header = [
            ('ETF', f"{pe.name} ({pe.id})"),
            ('Provider', p.name if p else pe.provider_id),
            ('Region', pe.region),
            ('Holdings date', ctx.etf_dates[etf_id]),
            ('Benchmark', f"{bm.name} ({bm.id}) as of {rows[0].benchmark_date}" if bm else "self (ETF's own holdings, market-cap weighted)"),
            ('Best ideas', f"top {big.MAX_BEST_IDEAS_PER_FUND} by delta = ETF weight - benchmark weight, delta > 0 and <= {big.HOLDING_DELTA_LIMIT_DROP_OFF:.0%}"),
        ]
        path = os.path.join(etf_dir, f"{_safe(pe.name or 'ETF')}_{pe.id}.xlsx")
        with pd.ExcelWriter(path, engine='openpyxl') as w:
            _write_sheet(w, 'Companies', company, {
                'Market value': MONEY, 'Market cap used': MONEY, 'ETF weight': PCT, 'Benchmark weight': PCT, 'Delta': PCT,
            }, header)
            _write_sheet(w, 'Listings', listings.sort_values(['Company', 'Listing']) if not listings.empty else listings,
                         {'Shares': NUM, 'Market value': MONEY, 'Provider weight': PCT})
            q = pd.DataFrame([{
                'Symbol': ctx.sym(qh.ticker_id), 'Name': ctx.tname(qh.ticker_id), 'Line id': h.id,
                'Shares': h.shares, 'Market value': h.market_value, 'Implied price': implied_price(h),
            } for qh in quarantined for h in qh.lines])
            _write_sheet(w, 'Quarantined', q if not q.empty else pd.DataFrame({'Note': [
                f"No ticker was listed on several lines with inconsistent prices (tolerance {DUP_PRICE_TOLERANCE})."
            ]}), {'Shares': NUM, 'Market value': MONEY, 'Implied price': NUM})
        written.append(path)
    return written


# ── Benchmark file ───────────────────────────────────────────────────────────

def _write_benchmark_file(ctx: _Context, out_dir: str) -> str:
    path = os.path.join(out_dir, 'benchmark.xlsx')
    used: set[str] = set()
    held = {a.ticker_id for a in ctx.analysis}
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        if ctx.mode == 'full_universe':
            snapshots = sorted({(a.benchmark_id, a.benchmark_date) for a in ctx.analysis if a.benchmark_id})
            for bm_id, bm_date in snapshots:
                bm = ctx.benchmarks[bm_id]
                rows = benchmark.fetch_latest_holdings_for_date(bm_id, bm_date)
                extra = [r.ticker_id for r in rows if r.ticker_id not in ctx.tickers]
                ctx.tickers.update({t.id: t for t in ticker.fetch_by_ids(extra)} if extra else {})
                df = pd.DataFrame([{
                    'Symbol': ctx.sym(r.ticker_id), 'Name': ctx.tname(r.ticker_id),
                    'Country': ctx.tickers[r.ticker_id].country if r.ticker_id in ctx.tickers else '',
                    'Region': ctx.tickers[r.ticker_id].region if r.ticker_id in ctx.tickers else '',
                    'Exchange': ctx.exchange(r.ticker_id),
                    'Market cap': r.market_cap, 'Weight': r.weight,
                    'Held by a fund ETF': 'Y' if r.ticker_id in held else '',
                } for r in rows]).sort_values('Weight', ascending=False)
                etfs = ", ".join(ctx.etf_label(a) for a in sorted({a.provider_etf_id for a in ctx.analysis if a.benchmark_id == bm_id}))
                _write_sheet(w, _sheet_name(bm.name, used), df, {'Market cap': MONEY, 'Weight': PCT}, [
                    ('Benchmark', f"{bm.name} ({bm.id})"), ('Region', bm.region), ('Snapshot date', bm_date),
                    ('Universe', f"{bm.cap_type} {bm.style_type}, USD market cap >= {bm.market_cap_min:,}, market-cap weighted, one row per company (its primary listing's cap), region by primary listing"),
                    ('Constituents', len(df)), ('Used for ETFs', etfs),
                ])
        else:
            for etf_id in sorted(ctx.etf_dates):
                rows = [a for a in ctx.analysis if a.provider_etf_id == etf_id and a.benchmark_weight is not None]
                df = pd.DataFrame([{
                    'Symbol': ctx.sym(a.ticker_id), 'Name': ctx.tname(a.ticker_id),
                    'Exchange': ctx.exchange(a.ticker_id),
                    'Market cap used': a.market_cap, 'Weight': a.benchmark_weight,
                } for a in rows]).sort_values('Weight', ascending=False)
                _write_sheet(w, _sheet_name(ctx.etfs[etf_id].name or str(etf_id), used), df, {'Market cap used': MONEY, 'Weight': PCT}, [
                    ('Benchmark', f"self — {ctx.etf_label(etf_id)} holdings, market-cap weighted"),
                    ('Holdings date', ctx.etf_dates[etf_id]),
                ])
    return path


# ── Best ideas file ──────────────────────────────────────────────────────────

def _eligibility(ctx: _Context, df: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    """(eligible, excluded_reason) per best-idea row: eligible in any of the fund's region/style
    buckets under the same masks and ranking cut the selection applies."""
    s = ctx.strategy
    styles = ['growth', 'value'] if (s.style.name == 'blend' and s.style.value is not None and s.style.growth is not None) else [s.style.name]
    ranking_level = s.ranking_to + (5 if s.allocation == 'market_cap' else 2)
    labels = {'etf': 'not one of the fund ETFs', 'style': 'style', 'cap': 'market cap', 'region': 'region (primary listing)',
              'exchange': 'exchange', 'esg': 'not ESG-qualified'}

    eligible = pd.Series(False, index=df.index)
    reasons: pd.Series = pd.Series([None] * len(df), index=df.index, dtype=object)
    for country_type in model_fund.strategy_country_types(s):
        for style in styles:
            masks = model_fund._eligibility_masks(df, s.provider_etfs or [], style, s.cap.name, country_type,
                                                  s.exchanges or [], s.esg_only, ctx.held_ids)
            ok = pd.Series(True, index=df.index)
            for m in masks.values():
                ok &= m
            eligible |= ok
            failed = pd.Series([
                ", ".join(labels[k] for k, m in masks.items() if not m.loc[i]) for i in df.index
            ], index=df.index)
            # keep the explanation with the fewest failed filters
            better = reasons.isna() | (failed.str.count(',') < reasons.fillna('').str.count(','))
            reasons = reasons.where(~(better & ~eligible), failed)
    rank_ok = (df['ranking'] <= ranking_level) & (df['ranking'] >= s.ranking_from)
    reasons = reasons.where(~eligible, None)
    reasons = reasons.where(~(eligible & ~rank_ok), f"rank outside {s.ranking_from}..{ranking_level}")
    return eligible & rank_ok, reasons


def _write_best_ideas_file(ctx: _Context, out_dir: str) -> str:
    path = os.path.join(out_dir, 'best_ideas.xlsx')
    df = ctx.ideas_df[ctx.ideas_df['provider_etf_id'].isin(set(ctx.etf_dates))].copy()
    stored = {(b.provider_etf_id, b.ticker_id): b for b in best_idea.fetch_for_etfs(ctx.etf_dates, ctx.mode)}
    eligible, reason = _eligibility(ctx, df)
    df['eligible'] = eligible
    df['excluded_reason'] = reason
    out = pd.DataFrame([{
        'Provider': ctx.provider_name(r['provider_etf_id']),
        'ETF': ctx.etf_label(r['provider_etf_id']),
        'Holdings date': r['value_date'],
        'Rank': r['ranking'],
        'Symbol': ctx.sym(r['ticker_id']),
        'Name': ctx.tname(r['ticker_id']),
        'Company (master)': ctx.sym(r['canonical_ticker_id']) if r['canonical_ticker_id'] != r['ticker_id'] else '',
        'ETF weight': stored[(r['provider_etf_id'], r['ticker_id'])].etf_weight if (r['provider_etf_id'], r['ticker_id']) in stored else None,
        'Benchmark weight': stored[(r['provider_etf_id'], r['ticker_id'])].benchmark_weight if (r['provider_etf_id'], r['ticker_id']) in stored else None,
        'Delta': r['delta'],
        'Style': r['style_type'],
        'Market cap': r['market_cap'],
        'Large cap': 'Y' if pd.notna(r['market_cap']) and r['market_cap'] >= LARGE_CAP_THRESHOLD else 'N',
        'Country': r['country'],
        'Region': r['region'],
        'Exchange': r['exchange'],
        'ESG': 'Y' if r['esg_qualified'] else 'N',
        'Eligible for fund': 'Y' if r['eligible'] else 'N',
        'Excluded because': r['excluded_reason'] or '',
    } for r in df.sort_values(['provider_etf_id', 'ranking']).to_dict('records')])
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        _write_sheet(w, 'Best ideas', out, {'ETF weight': PCT, 'Benchmark weight': PCT, 'Delta': PCT, 'Market cap': MONEY}, [
            ('Benchmark mode', ctx.mode),
            ('Per ETF', f"top {big.MAX_BEST_IDEAS_PER_FUND} companies by delta (0 < delta <= {big.HOLDING_DELTA_LIMIT_DROP_OFF:.0%})"),
            ('ETFs', len(ctx.etf_dates)),
        ])
    return path


# ── Fund file ────────────────────────────────────────────────────────────────

def _write_fund_file(ctx: _Context, out_dir: str) -> str:
    path = os.path.join(out_dir, 'fund.xlsx')
    s = ctx.strategy
    cand = {c.ticker_id: c for c in ctx.candidates}
    recomputed = {h.ticker_id: h for h in ctx.recomputed.holdings}

    def matches(h) -> bool:
        r = recomputed.get(h.ticker_id)
        return bool(r and r.ranking == h.ranking and r.source_etf_id == h.source_etf_id
                    and abs((r.weight or 0) - (h.weight or 0)) < 1e-9)

    def bucket(ticker_id: int) -> str:
        t = ctx.tickers.get(ticker_id)
        region = ('US' if t and t.region == 'US' else 'Non-US') if s.region and s.region.split else (s.region.name if s.region else 'all')
        return f"{region} / {t.style_type if t else '?'}"

    holdings = pd.DataFrame([{
        'Symbol': ctx.sym(h.ticker_id),
        'Name': ctx.tname(h.ticker_id),
        'Weight': h.weight,
        'Ranking': h.ranking,
        'Appearances': (ctx.buys[h.ticker_id].appearances if h.ticker_id in ctx.buys else (cand[h.ticker_id].appearances if h.ticker_id in cand else None)),
        'Max delta': h.max_delta,
        'Source ETF (top delta)': ctx.etf_label(h.source_etf_id),
        'All contributing ETFs': ", ".join(ctx.etf_label(e) for e in (
            (ctx.buys[h.ticker_id].all_provider_etf_ids if h.ticker_id in ctx.buys else None)
            or (cand[h.ticker_id].all_provider_ids if h.ticker_id in cand else []) or [])),
        'Market cap (weighting)': ctx.mc_map.get(h.ticker_id),
        'Bucket (region / style)': bucket(h.ticker_id),
        'Status': 'bought' if h.ticker_id in ctx.buys else 'kept from previous',
        'Matches stored': 'Y' if matches(h) else 'N',
    } for h in sorted(ctx.stored_holdings, key=lambda h: -(h.weight or 0))])

    selected = {h.ticker_id for h in ctx.stored_holdings}
    candidates = pd.DataFrame([{
        'Ranking': c.ranking,
        'Appearances': c.appearances,
        'Max delta': c.max_delta,
        'Symbol': ctx.sym(c.ticker_id),
        'Name': ctx.tname(c.ticker_id),
        'Bucket (region / style)': bucket(c.ticker_id),
        'Source ETF (top delta)': ctx.etf_label(c.source_etf_id),
        'All contributing ETFs': ", ".join(ctx.etf_label(e) for e in c.all_provider_ids),
        'Selected': 'Y' if c.ticker_id in selected else 'N',
        'Why not selected': '' if c.ticker_id in selected else (
            f"ranking beyond ranking_to ({s.ranking_to})" if c.ranking > s.ranking_to
            else "bucket / holding count already filled by better-ranked candidates"),
    } for c in sorted(ctx.candidates, key=lambda c: (c.ranking, -c.appearances, -c.max_delta))])

    strategy_rows = [(k, json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in s.model_dump().items()]
    summary = pd.DataFrame(
        [('Fund', f"{ctx.fund.name} ({ctx.fund.id})"), ('As of date', ctx.as_of_date), ('Inception date', ctx.inception),
         ('Holdings', len(ctx.stored_holdings)), ('Weights sum', sum(h.weight or 0 for h in ctx.stored_holdings)),
         ('All rows match stored', all(holdings['Matches stored'] == 'Y') if not holdings.empty else False),
         ('', '')] + [(f"strategy.{k}", v) for k, v in strategy_rows] + [('', '')] + [
            ('Candidate ranking level', s.ranking_to + (5 if s.allocation == 'market_cap' else 2)),
            ('Ranking = ', 'a company\'s best rank in any fund ETF on its latest date'),
            ('Order', 'ranking asc, appearances desc, max delta desc'),
            ('MC weight alpha / cap / floor', f"{model_fund.MC_WEIGHT_ALPHA} / {model_fund.MC_WEIGHT_CAP:.0%} / {model_fund.MC_WEIGHT_FLOOR:.0%}"),
        ],
        columns=['Item', 'Value'],
    )
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        _write_sheet(w, 'Holdings', holdings, {'Weight': PCT, 'Max delta': PCT, 'Market cap (weighting)': MONEY})
        _write_sheet(w, 'Candidates', candidates, {'Max delta': PCT})
        _write_sheet(w, 'Summary', summary)
    return path


# ── README ───────────────────────────────────────────────────────────────────

def _write_readme(ctx: _Context, out_dir: str, etf_files: list[str]) -> str:
    s = ctx.strategy
    bm_text = (
        "each ETF's configured external benchmark (see `benchmark.xlsx`): a synthetic large-cap, market-cap-weighted "
        "universe built from the FMP screener"
        if ctx.mode == 'full_universe' else
        "each ETF's own holdings, weighted by market cap (see `benchmark.xlsx`)"
    )
    lines = [
        f"# {ctx.fund.name} — holdings methodology",
        "",
        f"Fund id {ctx.fund.id} · as of {ctx.as_of_date}" + (" (inception)" if ctx.as_of_date == ctx.inception else f" (inception {ctx.inception})"),
        "",
        "## Strategy",
        "",
        "```json",
        json.dumps(ctx.fund.strategy, indent=2, default=str),
        "```",
        "",
        "## Steps",
        "",
        f"1. **ETF holdings** — `etfs/` has one workbook per constituent ETF ({len(etf_files)}), with the holdings file "
        f"downloaded from the provider for the date shown. A ticker listed on several lines is summed when all lines "
        f"imply the same price (within {DUP_PRICE_TOLERANCE - 1:.0%}), and excluded (quarantined) when they don't. "
        "Listings of one company (share classes such as GOOGL/GOOG, and foreign listings) are combined under the company's master ticker (*Master used*), measured by the company market cap of its primary listing.",
        f"2. **Benchmark** — mode `{ctx.mode}`: {bm_text}.",
        f"3. **Active weight** — for each company, ETF weight minus benchmark weight (delta). The top "
        f"{big.MAX_BEST_IDEAS_PER_FUND} companies with 0 < delta <= {big.HOLDING_DELTA_LIMIT_DROP_OFF:.0%} are the ETF's "
        "best ideas, ranked 1..N by delta. Every company and why it did or didn't qualify is in the ETF workbook (*Note*).",
        "4. **Best ideas** — `best_ideas.xlsx` lists all ETFs' best ideas, and whether each passes the fund's filters "
        f"(ETF list, style `{s.style.name}`, cap `{s.cap.name}`, region, exchange, ESG) and ranking cut.",
        "5. **Fund selection** — eligible ideas are combined per company: its best rank across the fund's ETFs, the "
        "number of ETFs naming it (appearances) and its largest delta. Companies are ordered by rank, then appearances, "
        f"then delta, and the top {s.holdings} are taken (split by region/style where the strategy says so). "
        f"Weights: {'market-cap based (square-root compressed, capped and floored)' if s.allocation == 'market_cap' else 'equal'}. "
        "See `fund.xlsx` — *Holdings* shows the justification per holding, *Candidates* shows who wasn't selected and why.",
        "",
        "## Constants",
        "",
        f"- Best ideas per ETF: {big.MAX_BEST_IDEAS_PER_FUND}; delta limit: {big.HOLDING_DELTA_LIMIT_DROP_OFF:.0%}",
        f"- Large-cap threshold: {LARGE_CAP_THRESHOLD:,}",
        f"- Duplicate-line price tolerance: {DUP_PRICE_TOLERANCE}",
        f"- Market-cap weighting: alpha {model_fund.MC_WEIGHT_ALPHA}, cap {model_fund.MC_WEIGHT_CAP:.0%}, floor {model_fund.MC_WEIGHT_FLOOR:.0%}",
        "",
        "## Caveat",
        "",
        "ETF, benchmark and best-idea numbers are the values stored when the fund was calculated. The fund-level filters "
        "(style, region, exchange, ESG) are evaluated against the tickers' *current* attributes, so for an old date a "
        "change since then can make the re-run selection differ — any such holding is flagged `Matches stored = N` in "
        "`fund.xlsx`.",
        "",
    ]
    path = os.path.join(out_dir, 'README.md')
    with open(path, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))
    return path


def run(fund_id: int, as_of_date: date | None = None) -> str:
    """Writes the methodology folder for fund_id (as of its inception unless as_of_date is
    given) and returns its path."""
    ctx = _Context(fund_id, as_of_date)
    out_dir = os.path.join(OUTPUT_ROOT, f"{_safe(ctx.fund.name)}_{fund_id}_{ctx.as_of_date}")
    os.makedirs(out_dir, exist_ok=True)

    etf_files = _write_etf_files(ctx, out_dir)
    _write_benchmark_file(ctx, out_dir)
    _write_best_ideas_file(ctx, out_dir)
    _write_fund_file(ctx, out_dir)
    _write_readme(ctx, out_dir, etf_files)
    return out_dir
